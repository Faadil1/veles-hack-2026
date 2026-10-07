"""Fill the official Veles Hack submission template and extend it into the pitch deck."""

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
IMG = HERE / "img"
DOCKER_IMAGE = sys.argv[1] if len(sys.argv) > 1 else "faadil12/hyperion:latest"
REPO = "github.com/Faadil1/veles-hack-2026"

NAVY = RGBColor(0x1F, 0x2A, 0x6B)
BLUE = RGBColor(0x1A, 0x1F, 0xD1)
INK = RGBColor(0x21, 0x21, 0x21)
MUTED = RGBColor(0x59, 0x59, 0x59)
CARD = RGBColor(0xF1, 0xF3, 0xFA)
GOOD = RGBColor(0x1E, 0x7A, 0x4F)
BAD = RGBColor(0xB3, 0x26, 0x1E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation(str(HERE / "template.pptx"))
L_TITLE_BODY = prs.slide_layouts.get_by_name("TITLE_AND_BODY")
L_TITLE_ONLY = prs.slide_layouts.get_by_name("TITLE_ONLY")


def set_text(shape, text):
    """Replace a placeholder's text but keep its first run's formatting."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    runs = p0.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p0.text = text
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)


def bullets(shape, items, size=14):
    tf = shape.text_frame
    first = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for i, item in enumerate(items):
        p = first if i == 0 else tf.add_paragraph()
        label, rest = item if isinstance(item, tuple) else ("", item)
        p.text = ""
        if label:
            r = p.add_run()
            r.text = label + " "
            r.font.bold = True
            r.font.size = Pt(size)
            r.font.color.rgb = NAVY
        r = p.add_run()
        r.text = rest
        r.font.size = Pt(size)
        r.font.color.rgb = INK
        p.space_after = Pt(8)


def box(slide, x, y, w, h, fill=CARD, name="card"):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    s.name = name
    return s


def text(slide, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         name="text"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        parts = para if isinstance(para, list) else [(para, {})] if isinstance(para, str) else [para]
        for part in parts:
            t, opt = part if isinstance(part, tuple) else (part, {})
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(opt.get("size", size))
            r.font.bold = opt.get("bold", bold)
            r.font.color.rgb = opt.get("color", color)
        p.space_after = Pt(opt.get("after", 4) if parts else 4)
    return tb


def picture(slide, path, x, y, w=None, h=None, crop=None, name="screenshot"):
    src = Path(path)
    if crop:
        im = Image.open(src).crop(crop)
        src = HERE / f"crop-{src.stem}-{crop[0]}-{crop[1]}.png"
        im.save(src)
    pic = slide.shapes.add_picture(str(src), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    pic.name = name
    pic.line.color.rgb = RGBColor(0xD0, 0xD4, 0xE4)
    pic.line.width = Pt(0.75)
    return pic


def new_slide(layout, title):
    s = prs.slides.add_slide(L_TITLE_BODY)  # same frame as the template's own content slides (no dots, numbered)
    set_text(s.shapes.title, title)
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx == 1:
            ph._element.getparent().remove(ph._element)
    return s


# 1 Title (template slide 1)
s1 = prs.slides[0]
set_text(s1.shapes.title, "Hyperion Steward")
sub = [ph for ph in s1.placeholders if ph.placeholder_format.idx == 1][0]
sub.text_frame.text = "Challenge 1 · HYPER-AI  |  says done only once it has seen it done"
for r in sub.text_frame.paragraphs[0].runs:
    r.font.size = Pt(16)
    r.font.color.rgb = WHITE
s1.notes_slide.notes_text_frame.text = (
    "Hyperion Steward is our Hyperion agent for the HYPER-AI IDE. One sentence: it only tells you something happened "
    "once it has seen it happen in your workspace.")

# 2 GitHub repo (template slide 2)
s2 = prs.slides[1]
body2 = [ph for ph in s2.placeholders if ph.placeholder_format.idx == 1][0]
bullets(body2, [("GitHub repo:", REPO),
                ("Docker image:", f"{DOCKER_IMAGE}  (exposes :8000/chat)"),
                ("Run:", f"docker run -p 8000:8000 --add-host host.docker.internal:host-gateway -e API_KEY=<key> "
                         f"{DOCKER_IMAGE}"),
                ("Licence:", "Apache-2.0, as in the official hyperion-starter"),
                ("Evidence:", "CI reports on the ci-evidence branch (official IDE images, digests recorded)")], size=14)
s2.notes_slide.notes_text_frame.text = "Repo, image, how to run it. Defaults match the starter: legion1, llama3.1, API_KEY."

# 3 Summary (template slide 3)
s3 = prs.slides[2]
body3 = [ph for ph in s3.placeholders if ph.placeholder_format.idx == 1][0]
body3.width = Inches(4.6)
bullets(body3, [
    "A drop-in /chat microservice for the HYPER-AI IDE: answers HYPER-AI questions from the official documents "
    "and turns plain language into IDE actions.",
    "We read the IDE the judges run, not only its tutorial. Its actions are fire-and-forget: a naive agent says "
    "done when nothing happened.",
    "Steward checks every profile with an exact copy of the IDE validator, then reads the workspace back and only "
    "claims what it saw. Deletes always ask; overwrites ask unless you named the file or Steward created it; "
    "file changes can be undone (a folder delete cannot).",
], size=13)
picture(s3, IMG / "03-ambiguous.png", 5.3, 1.45, w=4.3, crop=(1085, 625, 1595, 810), name="panel-screenshot")
text(s3, 5.3, 3.15, 4.3, 0.4, "The official IDE (organisers' GUI and backend images), a browser typing into the "
     "Hyperion panel in CI.", size=10, color=MUTED)
s3.notes_slide.notes_text_frame.text = (
    "What it is, what we found, what it does about it. The screenshot is the official IDE, run from the organisers' "
    "images in CI, with a browser typing into the Hyperion panel.")

# 4 Highlights (template slide 4): stat cards
s4 = prs.slides[3]
body4 = [ph for ph in s4.placeholders if ph.placeholder_format.idx == 1][0]
body4._element.getparent().remove(body4._element)
stats = [
    ("0", "disagreements with the IDE's own validator, over 5,536 test profiles", NAVY),
    ("9/9", "safety scenarios handled correctly (same tool calls, stub mirroring the IDE). Naive agent: 1/9", GOOD),
    ("0", "false \"done\" claims in the same scenarios. Naive agent: 5", GOOD),
    ("6/6", "live steps against the official backend image, every change verified by read-back", NAVY),
]
for i, (big, label, color) in enumerate(stats):
    x = 0.45 + i * 2.32
    box(s4, x, 1.35, 2.12, 2.6, name=f"stat-{i}")
    text(s4, x + 0.18, 1.55, 1.8, 0.9, big, size=44, color=color, bold=True, name=f"stat-{i}-value")
    text(s4, x + 0.18, 2.55, 1.8, 1.3, label, size=12, color=INK, name=f"stat-{i}-label")
text(s4, 0.45, 4.2, 9.1, 0.7,
     [[("Also: ", {"bold": True, "color": NAVY}),
       ("works without a model for the official example requests (create, delete, check, undo); plain-text replies "
        "because the panel does not render markdown; per-session memory sized to the provided 8k-token model.", {})]],
     size=12)
s4.notes_slide.notes_text_frame.text = (
    "Four numbers. Zero disagreements with the real validator. Nine out of nine safety scenarios against one out of "
    "nine for a naive agent, with identical tool calls. Zero false claims against five. Six of six steps live "
    "against the official backend image.")

# 5 What we found
s5 = new_slide(L_TITLE_ONLY, "What we found in the shipped IDE")
finds = [
    ("1", "Actions are fire-and-forget", "The GUI starts each action and never reports back. Failures only reach "
                                         "the status log."),
    ("2", "Many actions silently do nothing", "A name shared by two files, creating a file that exists, a closed "
                                              "IDE tab: no change, no signal. The tutorial says \"first match\"; "
                                              "the shipped GUI refuses instead."),
    ("3", "Nothing checks a profile unless asked", "Under YAML 1.2, isHighlyAvailable: no is a string and "
                                                   "schemaVersion: 1.1 a number. Both are saved as is."),
]
for i, (n, head, desc) in enumerate(finds):
    x = 0.45 + i * 3.08
    box(s5, x, 1.35, 2.88, 3.0, name=f"finding-{n}")
    circle = s5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.2), Inches(1.55), Inches(0.5), Inches(0.5))
    circle.fill.solid()
    circle.fill.fore_color.rgb = BLUE
    circle.line.fill.background()
    circle.name = f"finding-{n}-badge"
    ctf = circle.text_frame
    ctf.text = n
    ctf.paragraphs[0].alignment = PP_ALIGN.CENTER
    ctf.paragraphs[0].runs[0].font.size = Pt(16)
    ctf.paragraphs[0].runs[0].font.bold = True
    ctf.paragraphs[0].runs[0].font.color.rgb = WHITE
    text(s5, x + 0.2, 2.2, 2.5, 0.6, head, size=15, bold=True, color=NAVY, name=f"finding-{n}-head")
    text(s5, x + 0.2, 2.85, 2.5, 1.45, desc, size=12, name=f"finding-{n}-text")
text(s5, 0.45, 4.5, 9.1, 0.4, "Source: the GUI bundle and backend code inside donmichael/ide-gui:latest and "
                              "donmichael/ide-backend:latest.", size=10, color=MUTED)
s5.notes_slide.notes_text_frame.text = (
    "We ran the organisers' images and read their code. Three facts. One: actions are fire-and-forget. Two: plenty of "
    "them silently do nothing, and the tutorial even describes a behaviour the shipped GUI does not have. Three: "
    "nothing validates a profile unless someone asks. So a straightforward agent tells users it did things it did not "
    "do.")

# 6 How it works
s6 = new_slide(L_TITLE_ONLY, "How Steward handles every change")
steps = ["Scope guard", "Look up the target", "Check with the validator copy", "Ask when it matters",
         "Send the action", "Read the workspace back", "IDE validator, rollback if rejected"]
w, gap, y = 1.2, 0.12, 1.55
for i, label in enumerate(steps):
    x = 0.45 + i * (w + gap)
    fill = BLUE if i in (2, 5) else CARD
    b = box(s6, x, y, w, 1.25, fill=fill, name=f"step-{i + 1}")
    text(s6, x + 0.1, y + 0.12, w - 0.2, 1.0, [[(f"{i + 1}", {"bold": True, "size": 18,
                                                             "color": WHITE if i in (2, 5) else BLUE})],
                                              [(label, {"size": 11, "color": WHITE if i in (2, 5) else INK})]],
         name=f"step-{i + 1}-text")
text(s6, 0.45, 3.1, 4.4, 1.7, [
    [("The model handles language. ", {"bold": True, "color": NAVY}),
     ("It extracts intent and parameters; a builder emits the full profile, so it never hand-writes 40 typed "
      "fields.", {})],
    [("Every change goes through one gate. ", {"bold": True, "color": NAVY}),
     ("Whatever the model says, these steps run in this order.", {})]], size=12, name="how-left")
text(s6, 5.15, 3.1, 4.4, 1.7, [
    [("Answers are grounded. ", {"bold": True, "color": NAVY}),
     ("Retrieval over the official tutorial, cookbook and D3.3/D4.2/D4.3 deliverables, with citations.", {})],
    [("Receipts. ", {"bold": True, "color": NAVY}),
     ("/receipts/<user>/view shows each lookup, action, read-back, validator and model call.", {})]],
     size=12, name="how-right")
s6.notes_slide.notes_text_frame.text = (
    "One gate for every change. The two highlighted steps are what we added after reading the IDE: a validator copy "
    "proven identical to the real one, and a read-back that turns a fire-and-forget action into a verified one.")

# 7 Live demo
s7 = new_slide(L_TITLE_ONLY, "Live in the official IDE")
demo = IMG / "demo.png" if (IMG / "demo.png").exists() else IMG / "02-check-native.png"
picture(s7, demo, 0.45, 1.25, h=3.55, name="ide-screenshot")
text(s7, 6.95, 1.3, 2.65, 3.5, [
    [("Off-topic: ", {"bold": True, "color": NAVY}), ("declined before any model call.", {})],
    [("Check: ", {"bold": True, "color": NAVY}),
     ("the cookbook native example is valid but will not run (uvicorn in an nginx image).", {})],
    [("Ambiguous name: ", {"bold": True, "color": NAVY}), ("lists both files and asks.", {})],
    [("Create, fix, undo: ", {"bold": True, "color": NAVY}),
     ("each change is read back before Steward says it happened.", {})],
    [("Shown: ", {"bold": True, "color": NAVY}),
     ("the official create example, file opened in the editor. CI run with llama3.1 8B, 7/7.", {"color": MUTED})]],
     size=12, name="demo-notes")
s7.notes_slide.notes_text_frame.text = "Live demo in the official IDE (or this recording from CI)."

# 8 Criteria
s8 = new_slide(L_TITLE_ONLY, "The five criteria")
rows = [("Criterion", "Steward", "Proof"),
        ("1  /chat microservice: answers + NL to IDE actions", "SSE contract of the starter; builder; read-back",
         "CI: official images, browser-driven"),
        ("2  Guardrails", "Scope check before any model call", "Tests; off-topic scenario"),
        ("3  RAG on HYPER-AI docs", "Official tutorial, cookbook, deliverables; cited", "\"What is HyperAI?\" scenario"),
        ("4  Session memory", "Per user_id, sized to the 8k model", "Tests"),
        ("5  Human-in-the-loop (optional)", "Deletes always confirmed; overwrites unless path named or self-created; undo for file changes",
         "Live slice: undo restores identical bytes")]
tbl = s8.shapes.add_table(len(rows), 3, Inches(0.45), Inches(1.3), Inches(9.1), Inches(3.4)).table
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(3.3), Inches(3.2), Inches(2.6)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size = Pt(11 if r else 12)
        para.runs[0].font.bold = r == 0 or c == 0
        para.runs[0].font.name = FONT
        para.runs[0].font.color.rgb = WHITE if r == 0 else (NAVY if c == 0 else INK)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if r == 0 else (CARD if r % 2 else WHITE)
        cell.margin_left = cell.margin_right = Inches(0.08)
s8.notes_slide.notes_text_frame.text = "Each official criterion, how Steward meets it, and where it is proven."

# 9 Limits
s9 = new_slide(L_TITLE_ONLY, "What is proven, and what is not yet")
box(s9, 0.45, 1.3, 4.4, 2.9, name="proven")
text(s9, 0.65, 1.45, 4.0, 3.2, [
    [("Proven", {"bold": True, "color": GOOD, "size": 16})],
    "Validator parity with the real code (CI, official image)",
    "Create, 409 ambiguity, invalid profile stopped, delete, undo against the official backend image",
    "The official GUI in a browser: guardrail, check, ambiguity",
    "llama3.1 8B (the model family the organisers serve) in the official GUI, CI"], size=12, name="proven-text")
box(s9, 5.15, 1.3, 4.4, 2.9, name="not-yet")
text(s9, 5.35, 1.45, 4.0, 3.2, [
    [("Not yet", {"bold": True, "color": BAD, "size": 16})],
    "Answer quality on the organisers' server (legion1) not measured before submission",
    "Runnability findings are rules, not a deployment on HYPER-AI",
    "Steward cannot deploy; the IDE's Deploy button does",
    "AI-assisted development; full build provenance in the Autonomy Log"], size=12, name="not-yet-text")
s9.notes_slide.notes_text_frame.text = "What is proven and what is not. Every claim in the repo carries its evidence class."

prs.save(str(HERE / "Hyperion-Steward-VelesHack.pptx"))
print("saved")
